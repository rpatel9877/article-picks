"""End-to-end tests. Run with:  python -m pytest -q"""
import os
import sys
from datetime import datetime

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from config import Config  # noqa: E402

Config.ENABLE_SCHEDULER = False
Config.ANTHROPIC_API_KEY = ""
Config.SMTP_HOST = ""

import cycle as cyc  # noqa: E402
import db  # noqa: E402
import fetcher  # noqa: E402
import jobs  # noqa: E402
import scorer  # noqa: E402

TOPICS = ["dividend growth", "bond yields", "inflation", "interest rates", "earnings quality", "valuations",
          "private credit", "currency moves", "housing markets", "labor data"]
FINANCE_TEXT = " ".join(
    f"In section {n}, the author examines how {TOPICS[n % 10]} shaped portfolio returns during year {1970 + n}, "
    f"noting that patient investors who focused on {TOPICS[(n + 3) % 10]} saw steadier results than the market."
    for n in range(45))


class FakeResponse:
    def __init__(self, text="", status=200, ctype="text/html", url="https://example.org/a"):
        self.text, self.status_code, self.url = text, status, url
        self.headers = {"Content-Type": ctype}
        self.content = text.encode()


def page(body, title="Sample article", extra_head=""):
    paragraphs = "".join(f"<p>{s}.</p>" for s in body.split(".") if s.strip())
    return (f"<html><head><title>{title}</title>{extra_head}</head><body><article><h1>{title}</h1>"
            f"{paragraphs}</article></body></html>")


@pytest.fixture(autouse=True)
def isolated(tmp_path, monkeypatch):
    Config.DATABASE_PATH = str(tmp_path / "test.db")
    Config.OUTPUT_DIR = str(tmp_path / "reports")
    Config.OUTBOX_DIR = str(tmp_path / "outbox")
    db.init_db()
    db.save_settings({"team_members": "Priya Shah, priya@example.com\nMarcus Lee <marcus@example.com>\nTom Becker",
                      "report_recipients": "adam@example.com, megan@example.com"})

    def fake_get(url, **kwargs):
        if "paywalled" in url:
            return FakeResponse(page("Subscribe to continue reading this story. Already a subscriber? Sign in."), url=url)
        if "blocked" in url:
            return FakeResponse("Forbidden", status=403, url=url)
        if "short" in url:
            return FakeResponse(page("Stocks rose today on earnings. Investors cheered the market."), url=url)
        if "politics" in url:
            text = " ".join(f"Point {n}: Democrats and Republicans fight as Trump and Biden trade barbs, and "
                            f"conservatives blame liberals for economy problem number {n}." for n in range(40))
            return FakeResponse(page(text, "Partisan piece"), url=url)
        return FakeResponse(page(FINANCE_TEXT, "Why dividend growth wins"), url=url)

    monkeypatch.setattr(fetcher.requests, "get", fake_get)
    yield


# ---------- cycle logic ----------
def test_submissions_before_thursday_noon_count_for_this_week():
    tz = cyc.tz()
    monday = datetime(2026, 9, 21, 9, 0, tzinfo=tz)
    thursday_morning = datetime(2026, 9, 24, 11, 59, tzinfo=tz)
    thursday_after = datetime(2026, 9, 24, 12, 0, tzinfo=tz)
    assert cyc.open_cycle(monday) == "2026-09-24"
    assert cyc.open_cycle(thursday_morning) == "2026-09-24"
    assert cyc.open_cycle(thursday_after) == "2026-10-01"
    assert cyc.last_closed_cycle(thursday_after) == "2026-09-24"


def test_reminder_window_is_monday_morning_onward():
    tz = cyc.tz()
    assert not cyc.in_reminder_window(datetime(2026, 9, 21, 8, 0, tzinfo=tz))
    assert cyc.in_reminder_window(datetime(2026, 9, 21, 9, 0, tzinfo=tz))
    assert not cyc.in_reminder_window(datetime(2026, 9, 23, 9, 0, tzinfo=tz))


# ---------- URL handling ----------
def test_url_normalization_removes_tracking():
    a = fetcher.normalize_url("https://www.Example.com/story/?utm_source=x&id=5#top")
    b = fetcher.normalize_url("www.example.com/story?id=5")
    assert a == "https://www.example.com/story?id=5"
    assert b == a
    assert not fetcher.is_valid_url(fetcher.normalize_url("not a link"))


# ---------- paywall + scoring ----------
def test_known_paywall_domain_is_excluded():
    result = scorer.analyze("https://www.wsj.com/articles/something", db.get_settings())
    assert result["paywall"] == "likely"
    assert not result["eligible"]
    assert "Behind a paywall" in result["exclude_reasons"]


def test_paywall_phrases_detected():
    result = scorer.analyze("https://example.org/paywalled-story", db.get_settings())
    assert result["paywall"] == "likely"


def test_blocked_site_flagged_for_manual_check():
    result = scorer.analyze("https://example.org/blocked", db.get_settings())
    assert result["paywall"] == "unknown"
    assert not result["eligible"]


def test_good_article_is_eligible_and_scored():
    result = scorer.analyze("https://example.org/good-piece", db.get_settings())
    assert result["eligible"], result["exclude_reasons"]
    assert 0 < result["overall"] <= 100
    assert result["title"] == "Why dividend growth wins"


def test_short_and_partisan_articles_excluded():
    short = scorer.analyze("https://example.org/short-take", db.get_settings())
    partisan = scorer.analyze("https://example.org/politics", db.get_settings())
    assert not short["eligible"]
    assert not partisan["eligible"] and "Strong partisan angle" in partisan["exclude_reasons"]


def test_ai_json_parsing_is_robust():
    raw = 'Here you go:\n```json\n{"title":"T","scores":{"worldview_fit":"6","originality":2},"flags":{"partisan":1}}\n```'
    data = scorer._validate(scorer._extract_json(raw))
    assert data["scores"]["worldview_fit"] == 5 and data["scores"]["relevance"] == 3
    assert data["flags"]["partisan"] is True


# ---------- web app ----------
@pytest.fixture
def client():
    import app as app_module
    app = app_module.create_app(start_scheduler=False)
    app.config["TESTING"] = True
    return app.test_client()


def test_submit_flow_and_duplicates(client, monkeypatch):
    monkeypatch.setattr(jobs, "queue_analysis", lambda urls, refresh=False: None)
    resp = client.post("/", data={"submitter": "Priya Shah",
                                  "links": "https://example.org/good-piece?utm_source=x\nexample.org/other"})
    assert resp.status_code == 302
    subs = db.list_submissions(cyc.open_cycle())
    assert {s["url"] for s in subs} == {"https://example.org/good-piece", "https://example.org/other"}

    client.post("/", data={"submitter": "Priya Shah", "links": "https://example.org/good-piece"})
    assert len(db.list_submissions(cyc.open_cycle())) == 2  # duplicate ignored

    page_html = client.get("/").get_data(as_text=True)
    assert "Your links this week" in page_html and "2 links" in page_html


def test_submit_rejects_unknown_name_and_empty_links(client):
    assert client.post("/", data={"submitter": "Stranger", "links": "https://a.com/x"}).status_code == 400
    assert client.post("/", data={"submitter": "Tom Becker", "links": "   "}).status_code == 400


def test_admin_requires_login(client):
    assert client.get("/admin").status_code == 302
    client.post("/admin/login", data={"password": Config.ADMIN_PASSWORD})
    assert client.get("/admin").status_code == 200
    assert client.get("/admin/settings").status_code == 200


# ---------- the Thursday report ----------
def test_full_week_report(client, monkeypatch):
    from docx import Document

    week = cyc.open_cycle()
    urls = [f"https://example.org/good-piece-{i}" for i in range(7)] + [
        "https://www.ft.com/content/abc", "https://example.org/short-take", "https://example.org/politics"]
    people = ["Priya Shah", "Marcus Lee", "Tom Becker"]
    for n, url in enumerate(urls):
        db.add_submission(week, people[n % 3], url, "Worth a read" if n == 0 else "")
    db.add_submission(week, "Marcus Lee", urls[0])  # two people sent the same article

    path, items = jobs.build_report_file(week)
    assert os.path.exists(path)
    top = [i for i in items if i.get("rank")]
    assert len(top) == 5
    assert all(i["paywall"] != "likely" and i["eligible"] for i in top)
    assert next(i for i in items if i["url"] == urls[0])["submitters"] == ["Priya Shah", "Marcus Lee"]

    doc = Document(path)
    text = "\n".join(p.text for p in doc.paragraphs)
    for name in people:
        assert name in text
    for url in urls:
        assert url in text  # every link appears next to its submitter
    bold_urls = {r.text for p in doc.paragraphs for r in p.runs if r.bold and r.text.startswith("http")}
    assert bold_urls == {i["url"] for i in top}  # exactly the top 5 are bold

    # Emailing without SMTP saves to the outbox instead of failing
    sent, message = jobs.send_report(week)
    assert not sent and "outbox" in message
    assert os.listdir(Config.OUTBOX_DIR)


def test_tick_sends_each_job_once():
    tz = cyc.tz()
    week = "2026-09-24"
    db.add_submission(week, "Priya Shah", "https://example.org/good-piece")
    monday = datetime(2026, 9, 21, 9, 5, tzinfo=tz)
    thursday = datetime(2026, 9, 24, 12, 5, tzinfo=tz)

    assert jobs.tick(monday) == [("reminder", "saved")]
    assert jobs.tick(monday) == []  # never twice
    assert jobs.tick(thursday) == [("report", "saved")]
    assert jobs.tick(thursday) == []
    runs = {(r["job"], r["cycle"]) for r in db.recent_runs()}
    assert ("reminder", week) in runs and ("report", week) in runs


def test_ai_scoring_path_with_mocked_model(monkeypatch):
    import anthropic

    captured = {}

    class Block:
        type = "text"
        text = ('{"title":"Why dividend growth wins","summary":"S.","topic":"Dividends",'
                '"scores":{"worldview_fit":5,"originality":4,"relevance":5,"readability":4,"nonpartisan":5},'
                '"flags":{"too_technical":false,"too_short":false,"partisan":false,"off_topic":false},'
                '"why":"Fits our long-term view.","draft_intro":"A clear look at dividend growth."}')

    class FakeMessages:
        def create(self, **kwargs):
            captured.update(kwargs)
            return type("Msg", (), {"content": [Block()]})()

    class FakeClient:
        def __init__(self, **kwargs):
            self.messages = FakeMessages()

    monkeypatch.setattr(anthropic, "Anthropic", FakeClient)
    monkeypatch.setattr(Config, "ANTHROPIC_API_KEY", "test-key")
    db.save_settings({"worldview": "We are patient, long-term investors.", "past_examples": "Last week we shared..."})

    result = scorer.analyze("https://example.org/good-piece", db.get_settings())
    assert result["method"] == "ai" and result["eligible"]
    assert result["overall"] == round((0.30 * 5 + 0.25 * 4 + 0.20 * 5 + 0.15 * 4 + 0.10 * 5) / 5 * 100)
    prompt = captured["messages"][0]["content"]
    assert "patient, long-term investors" in prompt and "Last week we shared" in prompt
    assert "dividend growth" in prompt
