"""Builds the Thursday Word document: every link next to the person who sent it,
with the suggested top 5 in bold."""
from collections import defaultdict

from docx import Document
from docx.enum.text import WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

import cycle as cyc
from scorer import CRITERIA_LABELS, WEIGHTS

INK = RGBColor(0x17, 0x28, 0x2A)
MUTED = RGBColor(0x5B, 0x6B, 0x6A)
WARN = RGBColor(0xA4, 0x45, 0x2C)
GOOD = RGBColor(0x2E, 0x6A, 0x4E)
LINK_HEX = "1F4E79"
STAR = "\u2605"

PAYWALL_TEXT = {"free": "Free to read", "likely": "Paywall", "unknown": "Paywall not confirmed, please check"}


# ---------- low-level helpers ----------
def _hyperlink(paragraph, url, text, bold=False, size=None):
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    if bold:
        rpr.append(OxmlElement("w:b"))
    color = OxmlElement("w:color")
    color.set(qn("w:val"), LINK_HEX)
    rpr.append(color)
    if size:
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(int(size * 2)))
        rpr.append(sz)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def _para(doc, space_after=2, indent=None, keep_next=False):
    p = doc.add_paragraph()
    fmt = p.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(space_after)
    if indent is not None:
        fmt.left_indent = Inches(indent)
    fmt.keep_with_next = keep_next
    return p


def _run(p, text, bold=False, italic=False, color=None, size=None):
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if color is not None:
        r.font.color.rgb = color
    if size:
        r.font.size = Pt(size)
    return r


def _left_rule(p, hex_color="9A6B22"):
    ppr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    for k, v in (("w:val", "single"), ("w:sz", "18"), ("w:space", "8"), ("w:color", hex_color)):
        left.set(qn(k), v)
    borders.append(left)
    ppr.append(borders)


def _setup(doc):
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin"):
        setattr(section, side, Inches(0.9))
    section.top_margin = section.bottom_margin = Inches(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    for level, size, before in ((1, 15, 16), (2, 12, 10)):
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Georgia"
        # Theme font attributes override the named font in Word, so remove them
        rfonts = style.element.rPr.rFonts
        for attr in ("w:asciiTheme", "w:hAnsiTheme"):
            if rfonts.get(qn(attr)) is not None:
                del rfonts.attrib[qn(attr)]
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = INK
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.keep_with_next = True


def _words(n):
    if not n:
        return ""
    return f"about {round(n, -2):,} words" if n >= 1000 else f"about {max(round(n, -1), 10)} words"


def _meta_line(item):
    parts = [item.get("site") or item.get("domain")]
    if item.get("word_count"):
        parts.append(_words(item["word_count"]))
    if item.get("is_pdf"):
        parts.append("PDF")
    parts.append(PAYWALL_TEXT.get(item.get("paywall"), ""))
    if item.get("overall") is not None:
        parts.append(f"Score {item['overall']}/100")
    return "   |   ".join(p for p in parts if p)


def _names(names):
    return names[0] if len(names) == 1 else ", ".join(names[:-1]) + " and " + names[-1]


# ---------- sections ----------
def _top_entry(doc, item):
    p = _para(doc, space_after=1, keep_next=True)
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.first_line_indent = Inches(-0.3)
    _run(p, f"{item['rank']}.\t", bold=True, color=INK)
    _hyperlink(p, item["url"], item["title"], bold=True, size=11)

    p = _para(doc, space_after=1, indent=0.3, keep_next=True)
    _run(p, "Submitted by ", color=MUTED, size=9.5)
    _run(p, _names(item["submitters"]), bold=True, color=MUTED, size=9.5)
    _run(p, f"   |   {_meta_line(item)}", color=MUTED, size=9.5)

    if item.get("why"):
        p = _para(doc, space_after=8, indent=0.3)
        _run(p, item["why"], size=10)
    if item.get("paywall") == "unknown":
        p = _para(doc, space_after=8, indent=0.3)
        _run(p, f"Check before sending: {item.get('paywall_reason', 'paywall not confirmed')}.",
             italic=True, color=WARN, size=9.5)


def _list_entry(doc, item, submitter):
    top = bool(item.get("rank"))
    p = _para(doc, space_after=0, keep_next=True)
    if top:
        _run(p, f"{STAR} #{item['rank']}  ", bold=True, color=RGBColor(0x9A, 0x6B, 0x22))
    _hyperlink(p, item["url"], item["title"], bold=top, size=10.5)

    p = _para(doc, space_after=0, keep_next=True)
    _run(p, item["url"], bold=top, color=MUTED, size=8.5)

    p = _para(doc, space_after=0, keep_next=True)
    _run(p, _meta_line(item), bold=top, color=MUTED, size=9)

    others = [s for s in item["submitters"] if s != submitter]
    if others:
        p = _para(doc, space_after=0, keep_next=True)
        _run(p, f"Also sent by {_names(others)}", italic=True, color=MUTED, size=9)

    for who, note in item.get("notes", []):
        if who == submitter:
            p = _para(doc, space_after=0, keep_next=True)
            _run(p, f"{who.split()[0]}'s note: ", bold=True, size=9.5)
            _run(p, note, italic=True, size=9.5)

    p = _para(doc, space_after=9)
    if item.get("exclude_reasons"):
        _run(p, "Not suggested: " + "; ".join(item["exclude_reasons"]) + ". ", color=WARN, size=9.5)
    if item.get("why") and not top:
        _run(p, item["why"], size=9.5)
    elif top and item.get("summary"):
        _run(p, item["summary"], bold=True, size=9.5)


def build_report(path, cycle_id, items, settings):
    doc = Document()
    _setup(doc)

    total_links = len(items)
    submitters = sorted({s for i in items for s in i["submitters"]}, key=str.lower)
    top = sorted([i for i in items if i.get("rank")], key=lambda i: i["rank"])
    ai_used = any(i.get("method") == "ai" for i in items)

    # Title block
    p = _para(doc, space_after=2)
    _run(p, "Weekly Article Submissions", bold=True, color=INK, size=22).font.name = "Georgia"
    p = _para(doc, space_after=12)
    _run(p, f"For the Friday client email. Submissions closed {cyc.label(cycle_id)}.", color=MUTED, size=10)

    p = _para(doc, space_after=14)
    _left_rule(p)
    _run(p, f"{total_links} link{'s' if total_links != 1 else ''} from {len(submitters)} "
            f"team member{'s' if len(submitters) != 1 else ''}. ", bold=True)
    _run(p, f"The suggested top {len(top) or 5} are in bold and marked with a star in the full list. "
            "These are suggestions to speed up your review; the final pick is yours.")
    if len(top) < 5:
        _run(p, f" Only {len(top)} link{'s' if len(top) != 1 else ''} met every criterion this week.",
             color=WARN)
    if not ai_used:
        _run(p, " AI scoring was not available for this run, so links were ranked with basic rules.",
             italic=True, color=WARN)

    # Suggested top 5
    doc.add_heading(f"Suggested top {len(top) or 5}", level=1)
    if not top:
        p = _para(doc, space_after=8)
        _run(p, "No links met all of the criteria this week. See the full list below.", italic=True)
    for item in top:
        _top_entry(doc, item)

    # Full list grouped by submitter
    doc.add_heading("All submissions by team member", level=1)
    groups = defaultdict(list)
    for item in items:
        for name in item["submitters"]:
            groups[name].append(item)
    for name in sorted(groups, key=str.lower):
        links = sorted(groups[name], key=lambda i: (i.get("rank") or 99, -(i.get("overall") or -1)))
        doc.add_heading(f"{name}  ({len(links)} link{'s' if len(links) != 1 else ''})", level=2)
        for item in links:
            _list_entry(doc, item, name)

    # Draft intros (helps Megan with step 4)
    intros = [i for i in top if i.get("draft_intro")]
    if intros:
        doc.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)
        doc.add_heading("Draft intros for the top picks", level=1)
        p = _para(doc, space_after=8)
        _run(p, "Starting points for the Friday email. Edit freely before it goes to compliance.",
             italic=True, color=MUTED)
        for item in intros:
            p = _para(doc, space_after=2, keep_next=True)
            _run(p, item["title"], bold=True)
            p = _para(doc, space_after=10)
            _run(p, item["draft_intro"])

    # Method note
    doc.add_heading("How the suggestions were made", level=1)
    p = _para(doc, space_after=4)
    _run(p, "Every link was opened and read automatically. Links were left out of the suggestions if they "
            "were behind a paywall, strongly partisan, not about finance, investing, or economics, too "
            "technical, or too short. The rest were scored from 1 to 5 on each criterion below and combined "
            "into a score out of 100: ")
    _run(p, "; ".join(f"{CRITERIA_LABELS[k]} {int(w * 100)}%" for k, w in WEIGHTS.items()) + ".")
    p = _para(doc, space_after=4)
    _run(p, "Paywall checks use a list of subscription sites plus signals on the page itself. Sites that "
            "block automated reading are flagged for a manual check rather than assumed free.",
         color=MUTED, size=9.5)

    footer = doc.sections[0].footer.paragraphs[0]
    _run(footer, f"Generated automatically for the week ending {cyc.short_label(cycle_id)}. "
                 "Suggestions only.", color=MUTED, size=8)

    doc.save(path)
    return path
